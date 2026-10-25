"""The effects test workbook for a ported Gloomhaven class: what to play, turn by turn, what Gloomhaven shows for each
card half, and a column for what Frosthaven showed. Built from the pack itself, so it follows every re-port:

  - the class's cards (mods/gloomhaven/files/.../AbilityCards) and what each half plays (its AnimOverload, or the
    game's default for the ability: Attack, RangedAttack, Loot, Heal, PowerUp);
  - pack.json's animation table: the effects and the sound of each animator state;
  - tools/gh_tests/<Class>.json, hand-written: the test hand in play order ("turns"), a plain description of each
    state ("states"), what the last in-game run found ("status"), and the effects that are not on a card ("actions").
    Without it, the hand is picked to cover every state once and the descriptions are generated from the data.

    python tools/gh_testsheet.py build Brute        -> tests/Brute effects test.xlsx (an existing file is kept: --force)
    python tools/gh_testsheet.py read Brute         -> the rows whose result is not "Matches", with the notes

The workbook is the player's: tests/ is not published (the repo is a whitelist). Needs openpyxl.
"""

import argparse
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PACK = ROOT / "mods" / "gloomhaven"
SPECS = Path(__file__).resolve().parent / "gh_tests"
OUT = ROOT / "tests"

RESULTS = ["Matches", "Differs", "Missing", "Not seen"]
FONT = "Arial"

# ability kinds -> the animator state the game plays when the card names none
DEFAULT_STATE = {"Attack": "Attack", "Loot": "Loot", "Heal": "Heal", "Move": None}
KIND_NAMES = {"MergedMoveAttack": "Move and attack", "PreventDamage": "Prevent damage", "AddActiveBonus": "Active bonus",
              "Redirect": "Redirect attacks", "Kill": "Kill"}


# ---------------------------------------------------------------- the data

def strings():
    out = {}
    for line in (PACK / "strings.tsv").read_text(encoding="utf-8").splitlines():
        k, _, v = line.partition("\t")
        out[k] = v
    return out


def cards(cls):
    """[{key, name, level, initiative, halves: {Top: [ability], Bottom: [ability]}}] for the class, in card order.
    An ability: {kind, amount, range, area, jump, anim}."""
    names = strings()
    out = []
    for f in sorted((PACK / "files" / "Global" / "Characters" / cls / "AbilityCards").glob("*.yml")):
        text = f.read_text(encoding="utf-8")
        key = re.search(r"^Name:\s*\$ABILITY_CARD_(\w+)\$", text, re.M).group(1)
        card = {"key": key, "name": names.get("ABILITY_CARD_" + key, key),
                "level": re.search(r"^Level:\s*(\w+)", text, re.M).group(1),
                "initiative": int(re.search(r"^Initiative:\s*(\d+)", text, re.M).group(1)),
                "halves": {"Top": [], "Bottom": []}}
        half, ab = None, None
        for line in text.splitlines():
            m = re.match(r"^(\s*)([^\s:-][^:]*):\s*(.*?)\s*$", line)
            if not m:
                continue
            ind, k, v = len(m.group(1)), m.group(2), m.group(3)
            if ind == 0:
                half = None
            if ind == 4 and k in ("Top", "Bottom"):
                half = k
                continue
            if half is None:
                continue
            if ind == 12:
                ab = {"kind": None, "amount": None, "range": None, "area": False, "jump": False, "anim": None}
                card["halves"][half].append(ab)
            elif ab is not None and ind == 16:
                if k == "AnimOverload":
                    ab["anim"] = v
                elif k == "Range":
                    ab["range"] = v
                elif k == "AreaEffect":
                    ab["area"] = True
                elif k in ("ActiveBonus", "ActiveBonusDuration"):
                    ab["bonus"] = True
                elif ab["kind"] is None and k not in ("Filter", "Target", "Conditions", "ActiveBonus", "ActiveBonusDuration",
                                                      "SubAbilities", "MultiPassAttack", "XPPerTarget", "StatIsBasedOnX",
                                                      "AttackSourcesOnly", "StartAbilityRequirements", "Name", "Text",
                                                      "AreaEffectLayoutOverride", "PreviewEffectText", "AttackEffect",
                                                      "AbilityOverrides", "Elements"):
                    ab["kind"] = k
                    ab["amount"] = v or None
            elif ab is not None and ind > 16:
                if k == "Amount" and ab["amount"] is None:
                    ab["amount"] = v
                elif k == "Jump" and v.lower() == "true":
                    ab["jump"] = True
                elif k == "Move" and ab["kind"] == "MergedMoveAttack" and v:
                    ab["amount"] = v
        for h in card["halves"]:
            card["halves"][h] = [a for a in card["halves"][h] if a["kind"]]
        out.append(card)
    return out


def animation(cls):
    return json.loads((PACK / "pack.json").read_text(encoding="utf-8"))["model"][cls]["animation"]


def states_of(ab):
    """The animator states an ability plays, in order ("Movement" for a plain move)."""
    kind, anim = ab["kind"], ab["anim"]
    if anim == "None":
        return []
    if kind in ("Move",):
        return ["Jump" if ab["jump"] else "Movement"]
    if kind == "MergedMoveAttack":
        return ["Jump" if ab["jump"] else "Movement", "Attack"]
    if anim:
        return [anim]
    if kind == "Attack":
        return ["RangedAttack" if ab["range"] else "Attack"]
    return [DEFAULT_STATE.get(kind, "PowerUp")]


def plays(anim, state):
    """(clip, sound event, [effect names]) the plugin plays for <state>, from the pack's animation table."""
    for table in ("gh_effects", "lent_states"):
        pass
    if state in anim.get("gh_states", {}):
        fx = anim.get("gh_effects", {}).get(state) or {}
        return anim["gh_states"][state], anim.get("gh_audio", {}).get(state), _fx(fx)
    if state in anim.get("lent_states", {}):
        lent = anim["lent_states"][state]
        return lent.get("clip"), lent.get("audio"), _fx(lent.get("effects") or {})
    t = anim.get("states", {}).get("Base Layer." + state)
    if t is not None:
        return anim.get("clips", {}).get(anim.get("state_clips", {}).get(state)), t.get("audio"), _fx(t.get("effects") or {})
    return None, None, []


def _fx(play):
    names = [c["effect"] for c in play.get("caster", [])]
    for k in ("target", "area", "central", "aura"):
        names += [e["effect"] for e in (play.get(k) or {}).get("effects", [])]
    return names


def _pretty(name, cls):
    if not name:
        return ""
    n = re.sub(r"^(PlaySound_|P_|HE_)", "", name)
    n = re.sub(rf"^{cls}_?|_v\d+$", "", n)
    return re.sub(r"(?<=[a-z])(?=[A-Z])|_", " ", n).strip()


def generated(anim, state, cls):
    clip, audio, fx = plays(anim, state)
    parts = [f"Animation: {_pretty(clip, cls) or state}."]
    if fx:
        parts.append("Effects: " + ", ".join(_pretty(f, cls) for f in fx) + ".")
    if audio:
        parts.append(f"Sound: {_pretty(audio, cls)}.")
    return " ".join(parts)


def summary(abilities):
    out = []
    for a in abilities:
        name = KIND_NAMES.get(a["kind"], a["kind"])
        text = name + (f" {a['amount']}" if a["amount"] and a["amount"].lower() not in ("true",) else "")
        if a["range"]:
            text += f", range {a['range']}"
        if a["area"]:
            text += ", area"
        if a["jump"]:
            text += ", jump"
        if a.get("bonus"):
            text += " (bonus)"
        out.append(text)
    return "; ".join(out)


def pick_hand(all_cards):
    """Without a hand-written plan: cards picked greedily until every state any half plays is covered, ten at most."""
    want = {s for c in all_cards for h in c["halves"].values() for a in h for s in states_of(a)} - {"Movement"}
    hand, covered = [], set()
    while want - covered and len(hand) < 10:
        best = max(all_cards, key=lambda c: len({s for h in c["halves"].values() for a in h for s in states_of(a)} - covered))
        hand.append(best)
        covered |= {s for h in best["halves"].values() for a in h for s in states_of(a)}
    if len(hand) % 2:   # a turn needs two cards: the lowest-level card not yet taken
        hand.append(next(c for c in all_cards if c not in hand))
    return [{"top": hand[i]["key"], "bottom": hand[i + 1]["key"], "order": "", "setup": ""} for i in range(0, len(hand), 2)]


def version():
    v = (ROOT / "VERSION").read_text(encoding="utf-8").strip()
    try:
        sha = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    except OSError:
        sha = ""
    return f"{v} + {sha}" if sha else v


# ---------------------------------------------------------------- the workbook

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
RETEST = PatternFill("solid", fgColor="FCE4D6")


def _sheet(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.freeze_panes = "A2"


def _row(ws, values, inputs=()):
    ws.append(values)
    r = ws.max_row
    for i in range(1, len(values) + 1):
        c = ws.cell(r, i)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
        if i in inputs:
            c.fill = INPUT
    return r


def build(cls, force=False):
    spec_file = SPECS / f"{cls}.json"
    spec = json.loads(spec_file.read_text(encoding="utf-8")) if spec_file.exists() else {}
    all_cards = cards(cls)
    by_key = {c["key"].lower(): c for c in all_cards}
    anim = animation(cls)
    turns = spec.get("turns") or pick_hand(all_cards)
    notes, status = spec.get("states", {}), spec.get("status", {})
    built = version()

    def expected(state):
        return notes.get(state) or generated(anim, state, cls)

    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    plan = wb.create_sheet("Turn plan")
    tests = wb.create_sheet("Tests")

    # Tests: one row per card half of the hand, in play order, then the effects that are not on a card
    _sheet(tests, ["#", "Turn", "Card", "Half", "Ability", "Plays", "Expected (what Gloomhaven shows)",
                   "Status before this run", "Result in Frosthaven", "Notes"], [5, 6, 20, 8, 22, 18, 60, 26, 14, 40])
    seen, n = {}, 0
    for t, turn in enumerate(turns, 1):
        for half, key in (("Top", turn["top"]), ("Bottom", turn["bottom"])):
            card = by_key[key.lower()]
            abilities = card["halves"][half]
            states = [s for a in abilities for s in states_of(a) if s != "Movement"]
            n += 1
            lines, st = [], []
            for s in dict.fromkeys(states):
                if s in seen:
                    lines.append(f"Same {s} as row {seen[s]}.")
                else:
                    seen[s] = n
                    lines.append(expected(s))
                    if status.get(s):
                        st.append(status[s])
            if any(states_of(a) == ["Movement"] for a in abilities):
                lines.append("The move: see the Movement row.")
            r = _row(tests, [n, t, card["name"], half, summary(abilities), ", ".join(dict.fromkeys(states)) or "(move only)",
                             "\n".join(lines) or "Nothing of its own.", "\n".join(st), None, None], inputs=(9, 10))
            if any(x.startswith("RETEST") for x in st):
                tests.cell(r, 8).fill = RETEST
    for action in spec.get("actions") or [{"name": "Movement", "how": "Any move.", "expected": "Run cycle, footstep sounds and dust."}]:
        n += 1
        st = status.get(action["name"], "")
        r = _row(tests, [n, "", action["name"], "", action["how"], "", action["expected"], st, None, None], inputs=(9, 10))
        if st.startswith("RETEST"):
            tests.cell(r, 8).fill = RETEST
    last = tests.max_row
    tests.auto_filter.ref = f"A1:J{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(RESULTS) + '"', allow_blank=True,
                        promptTitle="Result", prompt="Matches, Differs, Missing (nothing played) or Not seen (didn't get to it)")
    tests.add_data_validation(dv)
    dv.add(f"I2:I{last}")
    for value, colour in (("Matches", "C6EFCE"), ("Differs", "FFC7CE"), ("Missing", "FFC7CE"), ("Not seen", "EDEDED")):
        tests.conditional_formatting.add(f"I2:I{last}", FormulaRule(formula=[f'$I2="{value}"'], fill=PatternFill("solid", fgColor=colour)))

    # Turn plan
    _sheet(plan, ["Turn", "Top", "Bottom", "Order", "Setup"], [6, 22, 22, 55, 40])
    for t, turn in enumerate(turns, 1):
        _row(plan, [t, by_key[turn["top"].lower()]["name"], by_key[turn["bottom"].lower()]["name"], turn.get("order", ""), turn.get("setup", "")])

    # How to use
    guide.column_dimensions["A"].width = 110
    lines = [
        (f"{cls}: effects test", Font(name=FONT, size=14, bold=True)),
        (f"Built {datetime.date.today():%d %b %Y} for mods {built}. Rebuilt from the pack when the class changes.", None),
        ("", None),
        ("1. Take the ten cards on the Turn plan sheet into a scenario with plenty of normal enemies, at a low scenario level.", None),
        ("2. Play the turns in order. Each Tests row is one card half; the moves are checked once, on the Movement row.", None),
        ("3. For each row, fill the yellow cells: Result (pick from the list) and Notes (what was different, louder or quieter, colours).", None),
        ("   Matches: looks and sounds like Gloomhaven.  Differs: something is off.  Missing: nothing played.  Not seen: didn't get to it.", None),
        ("4. 'Status before this run' says what the last run found. Rows marked RETEST had a fix since; the others only need a glance.", None),
        ("5. Save the file and tell Claude the results are in. Claude reads the sheet directly; no need to paste anything.", None),
        ("", None),
        ("Compare against Gloomhaven digital when unsure (Ctrl+F8 and Ctrl+F9 there unlock and max out the class).", None),
    ]
    for text, font in lines:
        guide.append([text])
        c = guide.cell(guide.max_row, 1)
        c.font = font or Font(name=FONT, size=11)
        c.alignment = Alignment(wrap_text=True)

    wb.active = wb.sheetnames.index("Tests")
    OUT.mkdir(exist_ok=True)
    path = OUT / f"{cls} effects test.xlsx"
    if path.exists() and not force:
        sys.exit(f"{path} exists (it may hold results); pass --force to replace it")
    wb.save(path)
    print(f"wrote {path}: {n} tests, {len(turns)} turns, for mods {built}")
    return path


def read(cls):
    path = OUT / f"{cls} effects test.xlsx"
    ws = load_workbook(path, data_only=True)["Tests"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    counts, open_rows = {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        result = row[col["Result in Frosthaven"]] or "(blank)"
        counts[result] = counts.get(result, 0) + 1
        if result != "Matches":
            open_rows.append(row)
    print(f"{path.name}: " + ", ".join(f"{k} {v}" for k, v in counts.items()))
    for row in open_rows:
        print(f"  #{row[col['#']]} {row[col['Card']]} {row[col['Half']] or ''}: {row[col['Result in Frosthaven']] or '(blank)'}"
              f" | {row[col['Notes']] or ''}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["build", "read"])
    ap.add_argument("cls")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    build(args.cls, args.force) if args.command == "build" else read(args.cls)
